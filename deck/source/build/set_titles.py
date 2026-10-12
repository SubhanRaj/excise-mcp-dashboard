"""Mark each slide's heading text box as the slide's title placeholder.

The slides are drawn on the Blank layout, so PowerPoint's outline view, the slide
navigator, accessibility readers and the PDF bookmarks see no title. This switches each
slide to the Title Only layout and tags its heading text box as the title.
Usage: python3 set_titles.py FILE.pptx [FILE.pptx ...]   (edits in place)
"""
import sys
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn


def heading(slide, first):
    """The title text box: 'Ask the Data' on the title slide, else the first text box."""
    for sh in slide.shapes:
        if sh.shape_type == 17 and sh.has_text_frame and sh.text_frame.text.strip():  # TEXT_BOX
            if not first or sh.text_frame.text.strip() == 'Ask the Data':
                return sh


def mark_titles(prs):
    title_only = next(l for l in prs.slide_layouts if l.name == 'Title Only')
    for n, s in enumerate(prs.slides, 1):
        sh = heading(s, n == 1)
        if sh is None or any(x.is_placeholder for x in s.shapes):
            continue
        rel = s.part.rels[next(k for k, r in s.part.rels.items() if r.reltype.endswith('/slideLayout'))]
        s.part._rels[rel.rId]._target = title_only.part
        nvpr = sh._element.find(qn('p:nvSpPr')).find(qn('p:nvPr'))
        for ph in nvpr.findall(qn('p:ph')):
            nvpr.remove(ph)
        nvpr.insert(0, etree.SubElement(nvpr, qn('p:ph'), type='ctrTitle' if n == 1 else 'title'))
        sh._element.find(qn('p:nvSpPr')).find(qn('p:cNvSpPr')).attrib.pop('txBox', None)
        sh.name = 'Title'


if __name__ == '__main__':
    for f in sys.argv[1:]:
        p = Presentation(f)
        mark_titles(p)
        p.save(f)
