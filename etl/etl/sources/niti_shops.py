"""The real-category Shops workbook -> shops / shop_years.
DATA_PIPELINE.md §Source adapters, §Shops.

Source is `Stats Dashboard Shop Types (2026-09-16)/5_Shops_All_Districts.xlsx`
(CLAUDE.md "Shop type reclassification for the stats dashboard"), not the
`Final Data for Submission` copy of the same sheet: that one only carries
NITI's 5 generic dropdown buckets (Retail/Composite/Model/WholeSale/Other),
78% of it literally "Other", which does not map onto any of excise_bank's 12
seeded license codes. This workbook carries UP's real 9 shop categories
instead, 7 of which map cleanly (§_CATEGORY_TO_CODE below); "Bar" and
"Wholesale Beer & Wine" still have no seeded code (FL6 / FL2B were never
added to `license_categories`) and load with `license_category_id = NULL`,
same as "Other" and "Retail" themselves - the row is still a real shop, just
uncategorized, never guessed onto a code that doesn't fit.

One sheet per district (the sheet name is the district, e.g. `BARA_BANKI`);
column B ("District") is not read for identity - `Bahraich`'s own sheet has a
whole FY's worth of rows with a wrong numeric value in that column (CLAUDE.md
"Bahraich" notes), so the file's own structure (one sheet per district) is
the reliable signal, not a free-text cell repeated on every row.

`shop_years` has one quantity slot (`mgq_bl`, bulk litres) where the source
has three (min IMFL / min country liquor / min beer, in mixed units - cases
for IMFL and beer, bulk litres for country liquor per the original DEO
instructions in CLAUDE.md). `_mgq_bl` fills it only when exactly one of the
three is present and already in bulk litres - never by guessing a
cases-to-litres conversion. `shop_years` also has no column at all for the
source's "lifted" (actual, as opposed to minimum) quantities, penalty
charged/recovered, or the gave-up/continued Yes-No pair beyond
`is_operational` (derived from "gave up shop": Yes -> not operational) - the
rest simply has nowhere to go in this schema and is not loaded.
"""

from collections.abc import Iterator
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

import asyncpg
import openpyxl

from etl import quarantine
from etl.normalize import (
    NormalizationError,
    parse_financial_year,
    resolve_district_id,
    seed_district_aliases,
)

START_HERE_SHEET = "start here"

# NITI's real UP shop category (as stored by CLAUDE.md's 2026-09-16
# reclassification) -> the one excise_bank license_categories.code it means.
# None = no seeded code exists for this category (reported in the run
# summary, never guessed at) or the category is itself one of the NITI
# bucket words passed through with no raw licence code left to recover.
_CATEGORY_TO_CODE: dict[str, str | None] = {
    "Country Liquor": "CL",
    "Foreign Liquor": "FL",
    "Beer": "BEER",
    "Composite": "FL5DB",
    "Model Shop": "MODEL",
    "Country Liquor with Beer": "CL5CC",
    "Premium Retail Vend": "FL4C",
    "Bar": None,
    "Wholesale Beer & Wine": None,
    "Other": None,
    "Retail": None,
}

# FLAGGED.md's 14 districts with an unresolved per-district Shops merge
# conflict (CLAUDE.md "Per-district Shops merge - COMPLETE"): their rows load
# (they're not dropped from the live stats dashboard either), but stay
# unpublished until a district confirms which side of the conflict is right.
FLAGGED_DISTRICTS = frozenset(
    {
        "Amroha",
        "Auraiya",
        "Azamgarh",
        "Bijnor",
        "Chandauli",
        "Farrukhabad",
        "Jaunpur",
        "Kasganj",
        "Kaushambi",
        "Mau",
        "Prayagraj",
        "Bhadohi",
        "Siddharthnagar",
        "Unnao",
    }
)

_VOLUME_UNIT_WORDS = ("bulk litres", "bl")


@dataclass(frozen=True)
class RunCounts:
    seen: int
    upserted: int
    quarantined: int


@dataclass(frozen=True)
class ShopRow:
    district_sheet: str
    shop_number: str
    seq: int
    financial_year_raw: str
    shop_type_raw: str | None
    allotment_method: str | None
    license_fee: Decimal | None
    min_imfl: Decimal | None
    min_country_liquor: Decimal | None
    min_beer: Decimal | None
    unit: str | None
    gave_up_shop: bool | None
    source_ref: str


def _num(value: object) -> Decimal | None:
    if value is None or value == "" or (isinstance(value, str) and value.strip() == "-"):
        return None  # '-' is this project's own "not collected" placeholder, same as blank
    if isinstance(value, int | float):
        return Decimal(str(value))
    if isinstance(value, str):
        # Indian-style comma grouping ('1,11,060') stored as literal text instead of
        # a clean number (CLAUDE.md's "Comma formatting" caveat). A handful of cells
        # still aren't a number after that (a cheque number, a stray character) -
        # treated as not-collected (None), same as a blank cell, rather than guessed.
        try:
            return Decimal(value.strip().replace(",", ""))
        except InvalidOperation:
            return None
    raise NormalizationError(f"unexpected cell type for a numeric column: {value!r}")


def _yesno(value: object) -> bool | None:
    if value is None:
        return None
    text = str(value).strip().lower()
    if text == "yes":
        return True
    if text == "no":
        return False
    return None


def read_shop_rows(path: str) -> Iterator[ShopRow]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet_name in workbook.sheetnames:
        if sheet_name.strip().lower() == START_HERE_SHEET:
            continue
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        next(rows)  # header
        # A generic/shared shop name can repeat within one district+FY (CLAUDE.md's
        # per-district-merge notes: name+FY is not a safe identity key on its own) -
        # disambiguate with an incrementing seq per (FY, shop_number) in row order.
        seen_in_fy: dict[str, int] = {}
        for i, r in enumerate(rows, start=2):
            if r[0] is None:
                continue
            fy_raw = str(r[2]).strip() if r[2] else ""
            shop_number = str(r[0]).strip()
            key = f"{fy_raw}\x00{shop_number}"
            seen_in_fy[key] = seen_in_fy.get(key, 0) + 1
            yield ShopRow(
                district_sheet=sheet_name,
                shop_number=shop_number,
                seq=seen_in_fy[key],
                financial_year_raw=fy_raw,
                shop_type_raw=str(r[3]).strip() if r[3] else None,
                allotment_method=str(r[4]).strip().lower() if r[4] else None,
                license_fee=_num(r[5]),
                min_imfl=_num(r[6]),
                min_country_liquor=_num(r[7]),
                min_beer=_num(r[8]),
                unit=str(r[12]).strip() if r[12] else None,
                gave_up_shop=_yesno(r[15]),
                source_ref=f"niti_shops:{path}:{sheet_name}:{i}",
            )
    workbook.close()


def _mgq_bl(row: ShopRow) -> Decimal | None:
    if row.unit is None or row.unit.strip().lower() not in _VOLUME_UNIT_WORDS:
        return None  # a Cases (or LPL) minimum needs a conversion factor not attempted here
    filled = [v for v in (row.min_imfl, row.min_country_liquor, row.min_beer) if v is not None]
    if len(filled) != 1:
        return None  # 0 or >1 simultaneously filled - which one is "the" MGQ is ambiguous
    return filled[0]


async def _resolve_financial_year_id(
    conn: asyncpg.Connection, cache: dict[str, int], raw: str
) -> int:
    if raw in cache:
        return cache[raw]
    start_year = parse_financial_year(raw)
    fy_id = await conn.fetchval("SELECT id FROM financial_years WHERE start_year = $1", start_year)
    if fy_id is None:
        raise NormalizationError(f"no financial_years row for start_year={start_year}")
    cache[raw] = int(fy_id)
    return cache[raw]


async def _load_license_category_ids(conn: asyncpg.Connection) -> dict[str, int]:
    rows = await conn.fetch("SELECT code, id FROM license_categories")
    return {r["code"]: int(r["id"]) for r in rows}


async def _lucknow_collisions(
    conn: asyncpg.Connection, lucknow_district_id: int, rows: list[ShopRow]
) -> set[str]:
    """Shop numbers this source assigns to Lucknow that already exist there from
    the live IESCMS dispatch import (a different numbering scheme, DATA_PIPELINE.md
    §Dispatches) - collisions get quarantined, never silently merged or duplicated
    (the task's own guardrail: this needs a human decision).

    `shops` has no source column, so a shop this same niti_shops source already
    inserted on an earlier run is told apart from a genuine IESCMS-origin shop by
    its `shop_years`: iescms_dispatch never writes shop_years at all (only
    `shops`), and no other source does either, so any shop_years row at all
    means this source already owns that shop - excluded from the collision
    set, or every re-run would flag its own previous rows as a collision with
    themselves.
    """
    existing = await conn.fetch(
        """
        SELECT s.shop_number FROM shops s
        WHERE s.district_id = $1
          AND NOT EXISTS (SELECT 1 FROM shop_years sy WHERE sy.shop_id = s.id)
        """,
        lucknow_district_id,
    )
    existing_numbers = {str(r["shop_number"]).strip().lower() for r in existing}
    incoming_numbers = {row.shop_number.lower() for row in rows if row.district_sheet == "LUCKNOW"}
    return existing_numbers & incoming_numbers


async def _upsert_shop(
    conn: asyncpg.Connection,
    *,
    district_id: int,
    shop_number: str,
    seq: int,
    license_category_id: int | None,
    published_at_now: bool,
) -> int:
    return int(
        await conn.fetchval(
            """
            INSERT INTO shops (district_id, shop_number, seq, display_name,
                                license_category_id, published_at)
            VALUES ($1, $2, $3, $4, $5, CASE WHEN $6 THEN now() ELSE NULL END)
            ON CONFLICT (district_id, shop_number, seq) DO UPDATE SET
                license_category_id = COALESCE(
                    EXCLUDED.license_category_id, shops.license_category_id
                ),
                published_at = COALESCE(shops.published_at,
                    CASE WHEN $6 THEN now() ELSE NULL END),
                updated_at = now()
            RETURNING id
            """,
            district_id,
            shop_number,
            seq,
            shop_number,
            license_category_id,
            published_at_now,
        )
    )


async def sync_rows(conn: asyncpg.Connection, run_id: int, rows: list[ShopRow]) -> RunCounts:
    seen = upserted = quarantined = 0
    fy_cache: dict[str, int] = {}
    # district_sheet -> (district_id, is_flagged); resolved once per sheet, not per row.
    district_cache: dict[str, tuple[int, bool]] = {}
    license_category_ids = await _load_license_category_ids(conn)
    unmapped_categories: dict[str, int] = {}

    lucknow_district_id = await resolve_district_id(conn, "Lucknow")
    collisions = await _lucknow_collisions(conn, lucknow_district_id, rows)

    for row in rows:
        seen += 1
        try:
            if row.district_sheet not in district_cache:
                district_key = row.district_sheet.replace("_", " ").title()
                resolved_id = await resolve_district_id(conn, district_key)
                canonical_name = await conn.fetchval(
                    "SELECT name FROM districts WHERE id = $1", resolved_id
                )
                district_cache[row.district_sheet] = (
                    resolved_id,
                    str(canonical_name) in FLAGGED_DISTRICTS,
                )
            district_id, is_flagged = district_cache[row.district_sheet]

            if row.district_sheet == "LUCKNOW" and row.shop_number.lower() in collisions:
                raise NormalizationError(
                    "shop_number collides with an existing IESCMS Lucknow shop: "
                    f"{row.shop_number!r}"
                )

            if row.shop_type_raw is None:
                raise NormalizationError("shop_type is required but the source cell was blank")

            fy_id = await _resolve_financial_year_id(conn, fy_cache, row.financial_year_raw)

            license_category_id = None
            if row.shop_type_raw in _CATEGORY_TO_CODE:
                code = _CATEGORY_TO_CODE[row.shop_type_raw]
                if code is not None:
                    license_category_id = license_category_ids[code]
                else:
                    unmapped_categories[row.shop_type_raw] = (
                        unmapped_categories.get(row.shop_type_raw, 0) + 1
                    )
            else:
                unmapped_categories[row.shop_type_raw] = (
                    unmapped_categories.get(row.shop_type_raw, 0) + 1
                )

            shop_id = await _upsert_shop(
                conn,
                district_id=district_id,
                shop_number=row.shop_number,
                seq=row.seq,
                license_category_id=license_category_id,
                published_at_now=not is_flagged,
            )

            await conn.execute(
                """
                INSERT INTO shop_years
                    (shop_id, financial_year_id, shop_type, mgq_bl,
                     license_fee_inr, settlement_mode, is_operational, source_ref)
                VALUES ($1, $2, $3, $4, $5, $6, $7, $8)
                ON CONFLICT (shop_id, financial_year_id) DO UPDATE SET
                    shop_type = EXCLUDED.shop_type, mgq_bl = EXCLUDED.mgq_bl,
                    license_fee_inr = EXCLUDED.license_fee_inr,
                    settlement_mode = EXCLUDED.settlement_mode,
                    is_operational = EXCLUDED.is_operational,
                    source_ref = EXCLUDED.source_ref
                """,
                shop_id,
                fy_id,
                row.shop_type_raw,
                _mgq_bl(row),
                row.license_fee,
                row.allotment_method,
                None if row.gave_up_shop is None else not row.gave_up_shop,
                row.source_ref,
            )
        except NormalizationError as exc:
            await quarantine.record(conn, run_id, row.__dict__, str(exc))
            quarantined += 1
            continue
        upserted += 1

    if unmapped_categories:
        # Not a per-row rejection (the shop and shop_year rows are inserted, just
        # with license_category_id left NULL) - one informational note per run,
        # not counted in `quarantined`, so the run summary can still say why.
        await quarantine.record(
            conn,
            run_id,
            {"unmapped_shop_categories": unmapped_categories},
            "shop_type values with no seeded license_categories code (FL6 'Bar', FL2B "
            "'Wholesale Beer & Wine', or a literal NITI passthrough word) - rows still "
            "loaded, license_category_id left NULL, not guessed",
        )

    return RunCounts(seen=seen, upserted=upserted, quarantined=quarantined)


async def sync(conn: asyncpg.Connection, run_id: int, path: str) -> RunCounts:
    await seed_district_aliases(conn)
    return await sync_rows(conn, run_id, list(read_shop_rows(path)))
