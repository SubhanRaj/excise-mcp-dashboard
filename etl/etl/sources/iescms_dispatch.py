"""IESCMS shop-wise wholesale-to-retail dispatch reports -> dispatches /
dispatch_strength_lines / shops. DATA_PIPELINE.md §Dispatches.

Two report layouts share this loader: the foreign-liquor-family report (one
cases/bottles/bulk-litres total per indent) and the country-liquor report
(the same totals broken down by liquor strength). read_fl_rows/read_cl_rows
turn each into a common DispatchRow, with strength_lines populated only for
country liquor. sync_rows does the shop upsert and the dispatches /
dispatch_strength_lines insert, kept separate from the xlsx reading so it's
testable against plain fixture rows with no openpyxl involved.
"""

from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import datetime

import asyncpg
import openpyxl

from etl import quarantine
from etl.normalize import NormalizationError, resolve_district_id

HEADER_ROW_INDEX = 15  # 0-indexed; rows above are the report's title/filter preamble


@dataclass(frozen=True)
class RunCounts:
    seen: int
    upserted: int
    quarantined: int


_CL_STRENGTH_LABELS = (
    "25% V/V",
    "36% V/V",
    "42.8% V/V 100 ML",
    "42.8% V/V 200 ML",
    "42.8% V/V",
    "28% V/V",
)


@dataclass(frozen=True)
class StrengthLine:
    strength_label: str
    requested_cases: float | None
    requested_bulk_litres: float | None
    dispatched_cases: float | None
    dispatched_bulk_litres: float | None


@dataclass(frozen=True)
class DispatchRow:
    district_name: str
    wholesale_license_type: str
    wholesale_license_number: str
    wholesale_entity_name: str
    retail_shop_ref: str  # the raw IESCMS shop id -> shops.shop_number
    retail_entity_name: str
    retail_license_type: str
    circle_sector: str | None
    indent_number: str
    indent_received_at: datetime | None
    indent_accepted_at: datetime | None
    transport_pass_issued_at: datetime | None
    tp_reference_no: str | None
    dispatched_bulk_litres: float
    duty_fee_inr: float
    source_ref: str
    requested_cases: float | None = None
    requested_bottles: float | None = None
    requested_bulk_litres: float | None = None
    dispatched_cases: float | None = None
    dispatched_bottles: float | None = None
    strength_lines: tuple[StrengthLine, ...] = field(default_factory=tuple)


def _num(value: object) -> float | None:
    if value is None or value == "":
        return None
    assert isinstance(value, int | float | str)  # openpyxl cells: numeric, or blank text
    return float(value)


def _dt(value: object) -> datetime | None:
    return value if isinstance(value, datetime) else None


def _load_sheet_rows(path: str) -> Iterator[tuple[int, tuple[object, ...]]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i <= HEADER_ROW_INDEX or row[0] is None:
            continue  # preamble, header, or a trailing blank/note line
        yield i, row
    workbook.close()


def read_fl_rows(path: str) -> Iterator[DispatchRow]:
    """WH-to-retail report for FL2-wholesaled shop types (FL5DB, FL4A, FL4C, ...)."""
    for i, row in _load_sheet_rows(path):
        yield DispatchRow(
            district_name=str(row[1]).strip(),
            wholesale_license_type=str(row[2]).strip(),
            wholesale_license_number=str(row[3]).strip(),
            wholesale_entity_name=str(row[4]).strip(),
            retail_shop_ref=str(row[6]).strip(),
            retail_entity_name=str(row[7]).strip(),
            retail_license_type=str(row[8]).strip(),
            circle_sector=str(row[10]).strip() if row[10] else None,
            indent_received_at=_dt(row[11]),
            indent_number=str(row[12]).strip(),
            requested_cases=_num(row[13]),
            requested_bottles=_num(row[14]),
            requested_bulk_litres=_num(row[15]),
            indent_accepted_at=_dt(row[16]),
            transport_pass_issued_at=_dt(row[17]),
            tp_reference_no=str(row[18]).strip() if row[18] else None,
            dispatched_cases=_num(row[19]),
            dispatched_bottles=_num(row[20]),
            dispatched_bulk_litres=_num(row[21]) or 0.0,
            duty_fee_inr=_num(row[22]) or 0.0,
            source_ref=f"iescms_dispatch_fl:{path}:{i}",
        )


def read_cl_rows(path: str) -> Iterator[DispatchRow]:
    """WH-to-retail report for CL2-wholesaled shop types (CL5C, CL5CC), broken
    down per liquor strength.
    """
    for i, row in _load_sheet_rows(path):
        strength_lines = tuple(
            StrengthLine(
                strength_label=label,
                requested_cases=_num(row[13 + k]),
                requested_bulk_litres=_num(row[19 + k]),
                dispatched_cases=_num(row[28 + k]),
                dispatched_bulk_litres=_num(row[34 + k]),
            )
            for k, label in enumerate(_CL_STRENGTH_LABELS)
            if any(
                v is not None
                for v in (
                    _num(row[13 + k]),
                    _num(row[19 + k]),
                    _num(row[28 + k]),
                    _num(row[34 + k]),
                )
            )
        )
        yield DispatchRow(
            district_name=str(row[1]).strip(),
            wholesale_license_type=str(row[2]).strip(),
            wholesale_license_number=str(row[3]).strip(),
            wholesale_entity_name=str(row[4]).strip(),
            retail_shop_ref=str(row[6]).strip(),
            retail_entity_name=str(row[7]).strip(),
            retail_license_type=str(row[8]).strip(),
            circle_sector=str(row[10]).strip() if row[10] else None,
            indent_received_at=_dt(row[11]),
            indent_number=str(row[12]).strip(),
            indent_accepted_at=_dt(row[25]),
            transport_pass_issued_at=_dt(row[26]),
            tp_reference_no=str(row[27]).strip() if row[27] else None,
            dispatched_bulk_litres=_num(row[40]) or 0.0,
            duty_fee_inr=_num(row[42]) or 0.0,
            source_ref=f"iescms_dispatch_cl:{path}:{i}",
            strength_lines=strength_lines,
        )


def _financial_year_start(dt: datetime) -> int:
    """India's fiscal year runs April-March: an August date's FY starts that
    same calendar year, a January date's FY started the year before.
    """
    return dt.year if dt.month >= 4 else dt.year - 1


async def _resolve_license_category_id(conn: asyncpg.Connection, code: str) -> int:
    category_id = await conn.fetchval("SELECT id FROM license_categories WHERE code = $1", code)
    if category_id is None:
        raise NormalizationError(f"unknown license category code: {code!r}")
    return int(category_id)


async def _resolve_financial_year_id(conn: asyncpg.Connection, dt: datetime | None) -> int:
    if dt is None:
        raise NormalizationError("no indent date to resolve a financial year from")
    start_year = _financial_year_start(dt)
    fy_id = await conn.fetchval("SELECT id FROM financial_years WHERE start_year = $1", start_year)
    if fy_id is None:
        raise NormalizationError(f"no financial_years row for start_year={start_year}")
    return int(fy_id)


async def _upsert_shop(
    conn: asyncpg.Connection, *, district_id: int, row: DispatchRow, license_category_id: int
) -> int:
    # The sheet gives one combined "Circle Sector" string (e.g. "Sector - 12"), not
    # separate circle/sector values, so it lands in sector_code alone.
    return int(
        await conn.fetchval(
            """
            INSERT INTO shops (district_id, shop_number, seq, display_name,
                                license_category_id, sector_code, published_at)
            VALUES ($1, $2, 1, $3, $4, $5, now())
            ON CONFLICT (district_id, shop_number, seq) DO UPDATE SET
                display_name = EXCLUDED.display_name,
                license_category_id = EXCLUDED.license_category_id,
                sector_code = EXCLUDED.sector_code,
                published_at = COALESCE(shops.published_at, now()),
                updated_at = now()
            RETURNING id
            """,
            district_id,
            row.retail_shop_ref,
            row.retail_entity_name,
            license_category_id,
            row.circle_sector,
        )
    )


async def sync_rows(conn: asyncpg.Connection, run_id: int, rows: list[DispatchRow]) -> RunCounts:
    seen = upserted = quarantined = 0

    for row in rows:
        seen += 1
        try:
            district_id = await resolve_district_id(conn, row.district_name)
            license_category_id = await _resolve_license_category_id(conn, row.retail_license_type)
            financial_year_id = await _resolve_financial_year_id(conn, row.indent_received_at)
            shop_id = await _upsert_shop(
                conn, district_id=district_id, row=row, license_category_id=license_category_id
            )
        except NormalizationError as exc:
            await quarantine.record(conn, run_id, row.__dict__, str(exc))
            quarantined += 1
            continue

        dispatch_id = await conn.fetchval(
            """
            INSERT INTO dispatches
                (district_id, financial_year_id, shop_id, wholesale_license_type,
                 wholesale_license_number, wholesale_entity_name, circle_sector,
                 indent_number, indent_received_at, indent_accepted_at,
                 transport_pass_issued_at, tp_reference_no, requested_cases,
                 requested_bottles, requested_bulk_litres, dispatched_cases,
                 dispatched_bottles, dispatched_bulk_litres, duty_fee_inr, source_ref,
                 published_at)
            VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10, $11, $12, $13, $14, $15,
                    $16, $17, $18, $19, $20, now())
            ON CONFLICT (indent_number) DO UPDATE SET
                published_at = COALESCE(dispatches.published_at, now())
            RETURNING id
            """,
            district_id,
            financial_year_id,
            shop_id,
            row.wholesale_license_type,
            row.wholesale_license_number,
            row.wholesale_entity_name,
            row.circle_sector,
            row.indent_number,
            row.indent_received_at,
            row.indent_accepted_at,
            row.transport_pass_issued_at,
            row.tp_reference_no,
            row.requested_cases,
            row.requested_bottles,
            row.requested_bulk_litres,
            row.dispatched_cases,
            row.dispatched_bottles,
            row.dispatched_bulk_litres,
            row.duty_fee_inr,
            row.source_ref,
        )
        for line in row.strength_lines:
            await conn.execute(
                """
                INSERT INTO dispatch_strength_lines
                    (dispatch_id, strength_label, requested_cases, requested_bulk_litres,
                     dispatched_cases, dispatched_bulk_litres)
                VALUES ($1, $2, $3, $4, $5, $6)
                ON CONFLICT (dispatch_id, strength_label) DO UPDATE SET
                    requested_cases = EXCLUDED.requested_cases,
                    requested_bulk_litres = EXCLUDED.requested_bulk_litres,
                    dispatched_cases = EXCLUDED.dispatched_cases,
                    dispatched_bulk_litres = EXCLUDED.dispatched_bulk_litres
                """,
                dispatch_id,
                line.strength_label,
                line.requested_cases,
                line.requested_bulk_litres,
                line.dispatched_cases,
                line.dispatched_bulk_litres,
            )
        upserted += 1

    return RunCounts(seen=seen, upserted=upserted, quarantined=quarantined)


async def sync(conn: asyncpg.Connection, run_id: int, path: str, report_kind: str) -> RunCounts:
    reader = read_fl_rows if report_kind == "fl" else read_cl_rows
    return await sync_rows(conn, run_id, list(reader(path)))
